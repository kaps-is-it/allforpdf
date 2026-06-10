import os
import io
import zipfile
import base64
import uuid
import html
import re
import aiofiles
from PIL import Image, ImageSequence
import img2pdf
import subprocess
import asyncio
import tempfile
import pathlib
import shutil
import markdown
import pandas as pd
from dataclasses import dataclass
from html.parser import HTMLParser
from pygments import highlight
from pygments.lexers import get_lexer_for_filename
from pygments.formatters import HtmlFormatter
from playwright.async_api import async_playwright
import fitz  # type: ignore[import-untyped]  # PyMuPDF — installed in Docker at runtime

from app.conversions.formats import (
    DOCUMENT_MEDIA_TYPES,
    IMAGE_CONVERSION_MAP,
    PDF_MEDIA_TYPES,
    get_document_download_name,
    get_pdf_download_name,
    normalize_image_extension,
)

try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except Exception:
    pass


@dataclass
class ConversionOutput:
    path: str
    filename: str
    media_type: str


async def _html_to_pdf(html_content, output_path, pdf_options=None):
    """Shared helper: renders HTML string to PDF using Playwright's Chromium."""
    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=['--no-sandbox', '--disable-setuid-sandbox']
        )
        page = await browser.new_page()
        await page.set_content(html_content, wait_until='networkidle')
        
        opts = {
            'path': output_path,
            'format': 'A4',
            'print_background': True
        }
        if pdf_options:
            opts.update(pdf_options)
        
        await page.pdf(**opts)
        await browser.close()


def _libreoffice_commands():
    return [
        'libreoffice',
        'soffice',
        'soffice.com',
        r'C:\Program Files\LibreOffice\program\soffice.com',
        r'C:\Program Files\LibreOffice\program\soffice.exe',
        r'C:\Program Files (x86)\LibreOffice\program\soffice.com',
        r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
    ]


def _run_libreoffice_convert(input_path, output_dir, target_ext):
    target_ext = target_ext if target_ext.startswith('.') else f'.{target_ext}'
    export_filters = {
        '.docx': 'docx:Office Open XML Text',
        '.odt': 'odt:writer8',
        '.rtf': 'rtf:Rich Text Format',
        '.txt': 'txt:Text',
    }
    target_name = export_filters.get(target_ext, target_ext.lstrip('.'))
    output_dir = pathlib.Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    input_path = pathlib.Path(input_path)

    profile_dir = tempfile.mkdtemp(prefix='lo_profile_')
    custom_env = os.environ.copy()
    custom_env['USERPROFILE'] = profile_dir
    custom_env['HOME'] = profile_dir

    last_error = None
    try:
        for cmd in _libreoffice_commands():
            try:
                result = subprocess.run(
                    [
                        cmd,
                        '--headless',
                        '--convert-to',
                        target_name,
                        '--outdir',
                        str(output_dir),
                        str(input_path),
                    ],
                    env=custom_env,
                    check=False,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                )
            except FileNotFoundError:
                continue

            expected_output = output_dir / f"{input_path.stem}{target_ext}"
            if result.returncode == 0 and expected_output.exists():
                return str(expected_output)

            candidates = sorted(
                output_dir.glob(f"{input_path.stem}.*"),
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
            for candidate in candidates:
                if candidate.suffix.lower() == target_ext:
                    return str(candidate)

            last_error = (result.stderr or result.stdout or '').strip()

        if last_error:
            raise Exception(f"LibreOffice failed during conversion: {last_error}")
        raise Exception("LibreOffice not found. Please install LibreOffice and ensure it is in the PATH for server deployment.")
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)


def _read_text_with_fallback(input_path):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return pathlib.Path(input_path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return pathlib.Path(input_path).read_text(encoding='utf-8', errors='replace')


def _plain_text_to_html(content, title='Converted Text'):
    paragraphs = []
    for block in re.split(r'\n\s*\n', content.strip()):
        escaped = html.escape(block).replace('\n', '<br />')
        paragraphs.append(f'<p>{escaped}</p>')

    body = '\n'.join(paragraphs) if paragraphs else '<p></p>'
    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>{html.escape(title)}</title>
  <style>
    body {{ font-family: Arial, sans-serif; line-height: 1.6; max-width: 760px; margin: 40px auto; color: #1f2937; }}
    p {{ margin: 0 0 1rem; white-space: normal; }}
  </style>
</head>
<body>
  {body}
</body>
</html>
"""


def _markdown_to_html_document(content, title='Converted Markdown'):
    html_body = markdown.markdown(content, extensions=[
        'pymdownx.superfences',
        'pymdownx.tasklist',
        'tables',
        'sane_lists',
    ])

    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>{html.escape(title)}</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif; line-height: 1.6; padding: 2rem; color: #24292f; }}
    code {{ background: rgba(175,184,193,0.2); padding: 0.2em 0.4em; border-radius: 6px; }}
    pre {{ padding: 16px; overflow: auto; background: #f6f8fa; border-radius: 6px; }}
    table {{ border-spacing: 0; border-collapse: collapse; width: 100%; margin-bottom: 16px; }}
    th, td {{ padding: 6px 13px; border: 1px solid #d0d7de; }}
  </style>
</head>
<body>
  {html_body}
</body>
</html>
"""


class _HTMLTextExtractor(HTMLParser):
    block_tags = {
        'address', 'article', 'aside', 'blockquote', 'div', 'footer', 'h1', 'h2',
        'h3', 'h4', 'h5', 'h6', 'header', 'li', 'main', 'p', 'section',
        'table', 'tr',
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style'}:
            self.skip_depth += 1
        elif tag == 'br':
            self.parts.append('\n')
        elif tag == 'li':
            self.parts.append('\n- ')
        elif tag in self.block_tags:
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in {'script', 'style'} and self.skip_depth:
            self.skip_depth -= 1
        elif tag in self.block_tags:
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.skip_depth:
            self.parts.append(data)

    def text(self):
        raw = ''.join(self.parts)
        raw = re.sub(r'[ \t\r\f\v]+', ' ', raw)
        raw = re.sub(r' *\n *', '\n', raw)
        raw = re.sub(r'\n{3,}', '\n\n', raw)
        return raw.strip() + '\n'


def _html_to_text(content):
    parser = _HTMLTextExtractor()
    parser.feed(content)
    return parser.text()


class _HTMLMarkdownExtractor(_HTMLTextExtractor):
    heading_tags = {'h1': '#', 'h2': '##', 'h3': '###', 'h4': '####', 'h5': '#####', 'h6': '######'}

    def handle_starttag(self, tag, attrs):
        if tag in self.heading_tags:
            self.parts.append(f"\n{self.heading_tags[tag]} ")
        elif tag in {'strong', 'b'}:
            self.parts.append('**')
        elif tag in {'em', 'i'}:
            self.parts.append('*')
        elif tag == 'a':
            self.parts.append('[')
        else:
            super().handle_starttag(tag, attrs)

    def handle_endtag(self, tag):
        if tag in self.heading_tags:
            self.parts.append('\n\n')
        elif tag in {'strong', 'b'}:
            self.parts.append('**')
        elif tag in {'em', 'i'}:
            self.parts.append('*')
        elif tag == 'a':
            self.parts.append(']')
        else:
            super().handle_endtag(tag)


def _html_to_markdown(content):
    parser = _HTMLMarkdownExtractor()
    parser.feed(content)
    text = parser.text()
    return text if text.strip() else '\n'


def _write_text_file(output_path, content):
    pathlib.Path(output_path).write_text(content, encoding='utf-8')
    return output_path


async def _convert_html_file_with_libreoffice(html_content, output_dir, target_ext):
    with tempfile.TemporaryDirectory(prefix='doc_html_') as temp_dir:
        html_path = pathlib.Path(temp_dir) / f'{uuid.uuid4()}.html'
        html_path.write_text(html_content, encoding='utf-8')
        return await asyncio.to_thread(_run_libreoffice_convert, str(html_path), output_dir, target_ext)

async def process_image_to_pdf(input_path, output_path):
    image = Image.open(input_path)
    pdf_bytes = img2pdf.convert(image.filename)
    async with aiofiles.open(output_path, "wb") as f:
        await f.write(pdf_bytes)
    return output_path


def _validate_image_conversion(input_ext, target_ext):
    input_ext = normalize_image_extension(input_ext)
    target_ext = normalize_image_extension(target_ext)
    allowed_targets = IMAGE_CONVERSION_MAP.get(input_ext)
    if not allowed_targets or target_ext not in allowed_targets:
        raise ValueError(f"Unsupported image conversion: {input_ext or 'unknown'} to {target_ext or 'unknown'}")
    return input_ext, target_ext


def _flatten_to_rgb(image, background=(255, 255, 255)):
    if image.mode in ('RGBA', 'LA') or ('transparency' in image.info):
        rgba = image.convert('RGBA')
        canvas = Image.new('RGBA', rgba.size, (*background, 255))
        canvas.alpha_composite(rgba)
        return canvas.convert('RGB')
    return image.convert('RGB')


def _load_image_frames(input_path):
    with Image.open(input_path) as image:
        frames = []
        durations = []
        for frame in ImageSequence.Iterator(image):
            frames.append(frame.copy())
            durations.append(frame.info.get('duration', image.info.get('duration', 100)) or 100)
        if not frames:
            frames.append(image.copy())
            durations.append(image.info.get('duration', 100) or 100)
    return frames, durations


def _save_frames_as_pdf(input_path, output_path):
    frames, _ = _load_image_frames(input_path)
    pdf_frames = [_flatten_to_rgb(frame) for frame in frames]
    pdf_frames[0].save(output_path, 'PDF', save_all=True, append_images=pdf_frames[1:])
    return output_path


def _save_icon(image, output_path):
    icon = Image.new('RGBA', (256, 256), (0, 0, 0, 0))
    source = image.convert('RGBA')
    source.thumbnail((256, 256), Image.Resampling.LANCZOS)
    icon.alpha_composite(source, ((256 - source.width) // 2, (256 - source.height) // 2))
    icon.save(
        output_path,
        format='ICO',
        sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)]
    )
    return output_path


def _save_embedded_svg(image, output_path):
    png_buffer = io.BytesIO()
    rgba = image.convert('RGBA')
    rgba.save(png_buffer, 'PNG')
    encoded_png = base64.b64encode(png_buffer.getvalue()).decode('ascii')
    width, height = rgba.size
    svg_content = (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}">'
        f'<image href="data:image/png;base64,{encoded_png}" width="{width}" height="{height}"/>'
        '</svg>'
    )
    pathlib.Path(output_path).write_text(svg_content, encoding='utf-8')
    return output_path


def _save_raster_image(input_path, output_path, input_ext, target_ext):
    frames, durations = _load_image_frames(input_path)

    if target_ext == '.jpg' or target_ext == '.jpeg':
        _flatten_to_rgb(frames[0]).save(output_path, 'JPEG', quality=95, optimize=True)
    elif target_ext == '.png':
        frames[0].convert('RGBA').save(output_path, 'PNG', optimize=True)
    elif target_ext == '.webp':
        webp_frames = [frame.convert('RGBA') for frame in frames]
        if input_ext == '.gif' and len(webp_frames) > 1:
            webp_frames[0].save(
                output_path,
                'WEBP',
                save_all=True,
                append_images=webp_frames[1:],
                duration=durations,
                loop=0,
                quality=90,
                method=6
            )
        else:
            webp_frames[0].save(output_path, 'WEBP', quality=90, method=6)
    elif target_ext == '.ico':
        _save_icon(frames[0], output_path)
    elif target_ext == '.bmp':
        _flatten_to_rgb(frames[0]).save(output_path, 'BMP')
    elif target_ext == '.gif':
        frames[0].convert('P', palette=Image.Palette.ADAPTIVE).save(output_path, 'GIF')
    elif target_ext == '.tiff':
        frames[0].convert('RGBA').save(output_path, 'TIFF')
    elif target_ext == '.svg':
        _save_embedded_svg(frames[0], output_path)
    else:
        raise ValueError(f"Unsupported raster image target: {target_ext}")

    return output_path


def _save_gif_frames_zip(input_path, output_path):
    frames, _ = _load_image_frames(input_path)
    stem = pathlib.Path(input_path).stem
    with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for index, frame in enumerate(frames, start=1):
            png_buffer = io.BytesIO()
            frame.convert('RGBA').save(png_buffer, 'PNG')
            zf.writestr(f"{stem}/frame_{index:04d}.png", png_buffer.getvalue())
    return output_path


def _save_gif_as_mp4(input_path, output_path):
    import imageio.v2 as imageio
    import numpy as np

    frames, durations = _load_image_frames(input_path)
    average_duration = max(sum(durations) / len(durations), 1)
    fps = max(1, min(60, round(1000 / average_duration)))

    normalized_frames = []
    max_width = max(frame.width for frame in frames)
    max_height = max(frame.height for frame in frames)
    if max_width % 2:
        max_width += 1
    if max_height % 2:
        max_height += 1

    for frame in frames:
        rgb = _flatten_to_rgb(frame)
        canvas = Image.new('RGB', (max_width, max_height), (255, 255, 255))
        canvas.paste(rgb, (0, 0))
        normalized_frames.append(np.asarray(canvas))

    with imageio.get_writer(output_path, fps=fps, codec='libx264', quality=8, macro_block_size=1) as writer:
        for frame in normalized_frames:
            writer.append_data(frame)
    return output_path


async def _render_svg_to_output(input_path, output_path, target_ext):
    async with aiofiles.open(input_path, 'rb') as f:
        svg_bytes = await f.read()

    encoded_svg = base64.b64encode(svg_bytes).decode('ascii')
    html_content = f"""
    <html>
    <head>
    <style>
        html, body {{
            margin: 0;
            padding: 0;
            background: transparent;
        }}
        img {{
            display: block;
        }}
    </style>
    </head>
    <body>
        <img id="svg-image" src="data:image/svg+xml;base64,{encoded_svg}" />
    </body>
    </html>
    """

    temp_png = None
    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=['--no-sandbox', '--disable-setuid-sandbox']
        )
        page = await browser.new_page()
        await page.set_content(html_content, wait_until='networkidle')
        dims = await page.evaluate("""
            () => {
                const img = document.getElementById('svg-image');
                const rect = img.getBoundingClientRect();
                return {
                    width: Math.ceil(img.naturalWidth || rect.width || 1024),
                    height: Math.ceil(img.naturalHeight || rect.height || 1024)
                };
            }
        """)
        width = max(1, min(int(dims['width']), 4096))
        height = max(1, min(int(dims['height']), 4096))
        await page.set_viewport_size({'width': width, 'height': height})
        await page.evaluate(
            """([width, height]) => {
                const img = document.getElementById('svg-image');
                img.style.width = `${width}px`;
                img.style.height = `${height}px`;
            }""",
            [width, height]
        )

        if target_ext == '.pdf':
            await page.pdf(
                path=output_path,
                width=f'{width}px',
                height=f'{height}px',
                print_background=True,
                margin={'top': '0', 'right': '0', 'bottom': '0', 'left': '0'}
            )
            await browser.close()
            return output_path

        fd, temp_png = tempfile.mkstemp(suffix='.png')
        os.close(fd)
        await page.locator('#svg-image').screenshot(
            path=temp_png,
            omit_background=(target_ext == '.png')
        )
        await browser.close()

    try:
        _save_raster_image(temp_png, output_path, '.png', target_ext)
    finally:
        if temp_png and os.path.exists(temp_png):
            os.remove(temp_png)

    return output_path


async def process_image_conversion(input_path, output_path, input_ext, target_ext):
    input_ext, target_ext = _validate_image_conversion(input_ext, target_ext)

    if input_ext == '.svg':
        return await _render_svg_to_output(input_path, output_path, target_ext)
    if target_ext == '.pdf':
        return await asyncio.to_thread(_save_frames_as_pdf, input_path, output_path)
    if input_ext == '.gif' and target_ext == '.zip':
        return await asyncio.to_thread(_save_gif_frames_zip, input_path, output_path)
    if input_ext == '.gif' and target_ext == '.mp4':
        return await asyncio.to_thread(_save_gif_as_mp4, input_path, output_path)

    return await asyncio.to_thread(_save_raster_image, input_path, output_path, input_ext, target_ext)

async def process_office_to_pdf(input_path, output_dir):
    # Uses LibreOffice headless for Word and PowerPoint presentations.
    return await asyncio.to_thread(_run_libreoffice_convert, input_path, output_dir, '.pdf')

async def process_spreadsheet_to_pdf(input_path, output_path, ext):
    def _read_data():
        if ext == '.csv':
            return pd.read_csv(input_path)
        elif ext == '.ods':
            return pd.read_excel(input_path, engine="odf")
        else:
            return pd.read_excel(input_path) # native xlsx
            
    df = await asyncio.to_thread(_read_data)
    
    # Generate HTML from the dataframe
    html_table = df.to_html(classes=['styled-table'], index=False)
    
    # Inject CSS to make tables look beautiful and fit to width
    html_content = f"""
    <html>
    <head>
    <style>
        @page {{ size: landscape; margin: 1cm; }}
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            font-size: 10px;
            color: #333;
        }}
        .styled-table {{
            border-collapse: collapse;
            width: 100%;
            table-layout: auto;
        }}
        .styled-table thead tr {{
            background-color: #8b5cf6;
            color: #ffffff;
            text-align: left;
        }}
        .styled-table th, .styled-table td {{
            padding: 8px 10px;
            border: 1px solid #dddddd;
            word-break: normal;
        }}
        .styled-table tbody tr:nth-of-type(even) {{
            background-color: #f3f3f3;
        }}
    </style>
    </head>
    <body>
        {html_table}
    </body>
    </html>
    """
    
    await _html_to_pdf(html_content, output_path, {
        'landscape': True
    })
    return output_path

async def process_ipynb_to_pdf(input_path, output_path):
    # Convert notebook to HTML first (avoids Windows asyncio subprocess bugs),
    # then render the HTML to PDF with our shared Playwright helper.
    html_path = str(input_path) + ".html"
    def _convert_to_html():
        subprocess.run(["jupyter", "nbconvert", "--to", "html", input_path, "--output", html_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    await asyncio.to_thread(_convert_to_html)
    
    async with aiofiles.open(html_path, 'r', encoding='utf-8') as f:
        html_content = await f.read()
    
    await _html_to_pdf(html_content, output_path, {
        'margin': {'top': '20px', 'right': '20px', 'bottom': '20px', 'left': '20px'}
    })
    
    try:
        os.remove(html_path)
    except:
        pass
        
    return output_path

async def process_text_code_to_pdf(input_path, output_path, ext):
    async with aiofiles.open(input_path, 'r', encoding='utf-8') as f:
        content = await f.read()

    html_content = ""
    # Use dark premium theme for code
    code_css = HtmlFormatter(style='monokai', full=False, cssclass="source").get_style_defs('.source')
    
    if ext == '.md':
        # Setup perfectly preserved Github Flavored Markdown
        html_body = markdown.markdown(content, extensions=[
            'pymdownx.superfences', 
            'pymdownx.tasklist', 
            'tables',
            'sane_lists'
        ])
        
        # Github theme inspired minimal CSS
        html_content = f"""
        <html><head><style>
        body{{ font-family: -apple-system,BlinkMacSystemFont,"Segoe UI",Helvetica,Arial,sans-serif; line-height: 1.6; padding: 2em; color: #24292f; }}
        h1, h2 {{ border-bottom: 1px solid #hsla(210,18%,87%,1); padding-bottom: .3em; }}
        code {{ background-color: rgba(175,184,193,0.2); padding: 0.2em 0.4em; border-radius: 6px; font-family: ui-monospace,SFMono-Regular,Menlo,Consolas,monospace; font-size: 85%;}}
        pre.source {{ padding: 16px; overflow: auto; line-height: 1.45; background-color: #272822; border-radius: 6px; color: #f8f8f2; }}
        {code_css}
        table {{ border-spacing: 0; border-collapse: collapse; width: 100%; margin-top: 0; margin-bottom: 16px; }}
        table th, table td {{ padding: 6px 13px; border: 1px solid #d0d7de; }}
        table tr:nth-child(2n) {{ background-color: #f6f8fa; }}
        ul.task-list {{ list-style-type: none; padding-left: 0; }}
        </style></head><body>{html_body}</body></html>
        """
    elif ext == '.html':
        html_content = content
    else:
        # Code formatting with Pygments
        try:
            lexer = get_lexer_for_filename(input_path)
            html_body = highlight(content, lexer, HtmlFormatter(style='monokai', cssclass="source", linenos=True))
        except Exception:
            # Fallback
            html_body = f"<pre class='source'><code>{content}</code></pre>"
            
        html_content = f"""
        <html><head><style>
        body{{ margin:0; font-family: ui-monospace, SFMono-Regular, Consolas, 'Courier New', monospace; background-color: #272822; color: #f8f8f2; }}
        .source {{ margin:0; padding: 20px; font-size: 12px; line-height: 1.5; }}
        .linenos {{ color: #75715e; margin-right: 15px; user-select: none; border-right: 1px solid #3e3d32; padding-right: 10px; }}
        {code_css}
        </style></head><body>{html_body}</body></html>
        """

    margin = (
        {'top': '40px', 'right': '40px', 'bottom': '40px', 'left': '40px'}
        if ext == '.md' else
        {'top': '0', 'right': '0', 'bottom': '0', 'left': '0'}
    )
    await _html_to_pdf(html_content, output_path, {'margin': margin})
    return output_path


async def process_document_conversion(input_path, output_dir, original_filename, input_ext, target_ext):
    output_dir = pathlib.Path(output_dir)
    output_path = output_dir / f"{uuid.uuid4()}{target_ext}"
    input_ext = '.html' if input_ext == '.htm' else input_ext

    if target_ext == '.pdf':
        if input_ext == '.txt':
            content = _read_text_with_fallback(input_path)
            await _html_to_pdf(
                _plain_text_to_html(content, pathlib.Path(original_filename).stem),
                str(output_path),
                {'margin': {'top': '40px', 'right': '40px', 'bottom': '40px', 'left': '40px'}},
            )
        elif input_ext == '.md':
            content = _read_text_with_fallback(input_path)
            await _html_to_pdf(
                _markdown_to_html_document(content, pathlib.Path(original_filename).stem),
                str(output_path),
                {'margin': {'top': '40px', 'right': '40px', 'bottom': '40px', 'left': '40px'}},
            )
        elif input_ext == '.html':
            await process_text_code_to_pdf(input_path, str(output_path), '.html')
        else:
            converted = await asyncio.to_thread(_run_libreoffice_convert, input_path, str(output_dir), '.pdf')
            shutil.move(converted, output_path)
    elif target_ext == '.html':
        if input_ext == '.txt':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, _plain_text_to_html(content, pathlib.Path(original_filename).stem))
        elif input_ext == '.md':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, _markdown_to_html_document(content, pathlib.Path(original_filename).stem))
        else:
            converted = await asyncio.to_thread(_run_libreoffice_convert, input_path, str(output_dir), '.html')
            shutil.move(converted, output_path)
    elif target_ext == '.txt':
        if input_ext == '.html':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, _html_to_text(content))
        elif input_ext == '.md':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, _html_to_text(_markdown_to_html_document(content)))
        else:
            converted = await asyncio.to_thread(_run_libreoffice_convert, input_path, str(output_dir), '.txt')
            shutil.move(converted, output_path)
    elif target_ext == '.md':
        if input_ext == '.txt':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, content if content.endswith('\n') else f'{content}\n')
        elif input_ext == '.html':
            content = _read_text_with_fallback(input_path)
            _write_text_file(output_path, _html_to_markdown(content))
        else:
            with tempfile.TemporaryDirectory(prefix='doc_md_') as temp_dir:
                converted_html = await asyncio.to_thread(_run_libreoffice_convert, input_path, temp_dir, '.html')
                html_content = _read_text_with_fallback(converted_html)
                _write_text_file(output_path, _html_to_markdown(html_content))
    elif target_ext in {'.docx', '.odt', '.rtf'}:
        if input_ext == '.txt':
            content = _read_text_with_fallback(input_path)
            converted = await _convert_html_file_with_libreoffice(
                _plain_text_to_html(content, pathlib.Path(original_filename).stem),
                str(output_dir),
                target_ext,
            )
        elif input_ext == '.md':
            content = _read_text_with_fallback(input_path)
            converted = await _convert_html_file_with_libreoffice(
                _markdown_to_html_document(content, pathlib.Path(original_filename).stem),
                str(output_dir),
                target_ext,
            )
        else:
            converted = await asyncio.to_thread(_run_libreoffice_convert, input_path, str(output_dir), target_ext)
        shutil.move(converted, output_path)
    else:
        raise ValueError(f"Unsupported document conversion target: {target_ext}")

    return ConversionOutput(
        path=str(output_path),
        filename=get_document_download_name(original_filename, target_ext),
        media_type=DOCUMENT_MEDIA_TYPES[target_ext],
    )


async def process_document_to_images(input_path, output_zip_path, original_filename):
    """
    Converts a .pptx or .pdf file into individual PNG page images, then bundles
    them into a single ZIP file. Returns the path to the generated ZIP.
    
    Steps:
        - PPTX: Use LibreOffice headless to convert PPTX → PDF first
        - PDF: Use the file directly
        - Then: Use PyMuPDF (fitz) to rasterize each PDF page as a high-res PNG
        - Finally: Package all PNGs into a ZIP file
    """
    
    input_ext = os.path.splitext(input_path)[1].lower()
    temp_dir = tempfile.mkdtemp(prefix='doc_to_img_')
    
    try:
        if input_ext == '.pdf':
            # PDF input: use directly, no conversion needed
            pdf_path = input_path
        else:
            # PPTX input: convert to PDF via LibreOffice first
            pdf_output_dir = os.path.join(temp_dir, 'pdf')
            os.makedirs(pdf_output_dir, exist_ok=True)
            
            def _convert_to_pdf():
                profile_dir = tempfile.mkdtemp(prefix='lo_profile_')
                custom_env = os.environ.copy()
                custom_env['USERPROFILE'] = profile_dir
                custom_env['HOME'] = profile_dir
                
                cmds = [
                    'libreoffice', 
                    'soffice',
                    'soffice.com',
                    r'C:\Program Files\LibreOffice\program\soffice.com',
                    r'C:\Program Files\LibreOffice\program\soffice.exe',
                    r'C:\Program Files (x86)\LibreOffice\program\soffice.com',
                    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe'
                ]
                success = False
                try:
                    for cmd in cmds:
                        try:
                            subprocess.run([
                                cmd, "--headless", "--convert-to", "pdf",
                                "--outdir", pdf_output_dir, input_path
                            ], env=custom_env, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                            success = True
                            break
                        except FileNotFoundError:
                            continue
                        except subprocess.CalledProcessError:
                            raise Exception("LibreOffice failed during PPTX-to-PDF conversion.")
                    if not success:
                        raise Exception("LibreOffice not found. Required for PPTX conversion.")
                finally:
                    shutil.rmtree(profile_dir, ignore_errors=True)
            
            await asyncio.to_thread(_convert_to_pdf)
            
            # Find the generated PDF
            input_basename = os.path.splitext(os.path.basename(input_path))[0]
            pdf_path = os.path.join(pdf_output_dir, f"{input_basename}.pdf")
            if not os.path.exists(pdf_path):
                raise Exception("PDF generation from PPTX failed — no output file created.")
        
        # Step 2: Rasterize PDF pages into PNG images using PyMuPDF at 300 DPI
        original_stem = os.path.splitext(original_filename)[0]
        
        def _render_and_zip():
            # Suppress noisy MuPDF warnings about broken structure trees
            # (common in LibreOffice-generated PDFs, harmless but spammy)
            fitz.TOOLS.mupdf_display_errors(False)
            fitz.TOOLS.mupdf_display_warnings(False)
            
            doc = fitz.open(pdf_path)
            if doc.page_count == 0:
                doc.close()
                raise Exception("No slides found in the presentation.")
            
            # Strip the broken StructTreeRoot from the PDF catalog to prevent
            # "No common ancestor in structure tree" errors and improve performance
            try:
                cat = doc.pdf_catalog()
                doc.xref_set_key(cat, "StructTreeRoot", "null")
            except Exception:
                pass  # Not critical — proceed even if stripping fails
            
            # 300 DPI = 300/72 ≈ 4.17x zoom factor
            zoom = 300 / 72
            matrix = fitz.Matrix(zoom, zoom)
            
            with zipfile.ZipFile(output_zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
                for i in range(doc.page_count):
                    page = doc[i]
                    pix = page.get_pixmap(matrix=matrix, alpha=False)
                    png_data = pix.tobytes("png")
                    # Store as: OriginalName/Slide_1.png, OriginalName/Slide_2.png, etc.
                    arcname = f"{original_stem}/Slide_{i + 1}.png"
                    zf.writestr(arcname, png_data)
            
            doc.close()
        
        await asyncio.to_thread(_render_and_zip)
        
        return output_zip_path
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def _get_pdf_text_pages(input_path):
    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")
        return [page.get_text("text").strip() for page in doc]


def _write_pdf_text_file(input_path, output_path):
    pages = _get_pdf_text_pages(input_path)
    content = []
    for index, text in enumerate(pages, start=1):
        if len(pages) > 1:
            content.append(f"--- Page {index} ---")
        content.append(text)
    pathlib.Path(output_path).write_text("\n\n".join(content).strip() + "\n", encoding="utf-8")
    return output_path


def _write_pdf_html_file(input_path, output_path):
    page_html = []
    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")
        for index, page in enumerate(doc, start=1):
            page_html.append(
                f'<section class="pdf-page" aria-label="Page {index}">'
                f"{page.get_text('html')}"
                "</section>"
            )

    html = f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>Converted PDF</title>
  <style>
    body {{ margin: 0; padding: 24px; background: #f6f7fb; font-family: Arial, sans-serif; }}
    .pdf-page {{ position: relative; margin: 0 auto 24px; background: white; box-shadow: 0 8px 24px rgba(15, 23, 42, 0.14); overflow: hidden; }}
  </style>
</head>
<body>
  {''.join(page_html)}
</body>
</html>
"""
    pathlib.Path(output_path).write_text(html, encoding="utf-8")
    return output_path


def _write_pdf_docx(input_path, output_path):
    try:
        from docx import Document
    except ImportError as exc:
        raise Exception("python-docx is required for PDF to DOCX conversion.") from exc

    pages = _get_pdf_text_pages(input_path)
    document = Document()
    document.core_properties.title = "Converted PDF"

    for index, text in enumerate(pages, start=1):
        if index > 1:
            document.add_page_break()
        if len(pages) > 1:
            document.add_heading(f"Page {index}", level=1)
        lines = text.splitlines() or ["No selectable text was found on this page."]
        for line in lines:
            document.add_paragraph(line)

    document.save(output_path)
    return output_path


def _write_pdf_pptx(input_path, output_path):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise Exception("python-pptx is required for PDF to PPTX conversion.") from exc

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]

    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")

        first_rect = doc[0].rect
        prs.slide_width = int(first_rect.width / 72 * 914400)
        prs.slide_height = int(first_rect.height / 72 * 914400)

        matrix = fitz.Matrix(2, 2)
        for page in doc:
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            image_stream = io.BytesIO(pix.tobytes("png"))
            image_stream.seek(0)
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                image_stream,
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

    prs.save(output_path)
    return output_path


def _write_pdf_xlsx(input_path, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    pages = _get_pdf_text_pages(input_path)
    workbook = Workbook()

    for page_index, text in enumerate(pages, start=1):
        worksheet = workbook.active if page_index == 1 else workbook.create_sheet()
        worksheet.title = f"Page {page_index}"
        worksheet["A1"] = f"Page {page_index}"
        worksheet["A1"].font = Font(bold=True)

        lines = text.splitlines() or ["No selectable text was found on this page."]
        for row_index, line in enumerate(lines, start=2):
            cells = line.split("\t") if "\t" in line else [line]
            for column_index, value in enumerate(cells, start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 80)

    workbook.save(output_path)
    return output_path


def _render_pdf_page_bytes(page, target_ext):
    if target_ext == ".svg":
        return page.get_svg_image().encode("utf-8")

    matrix = fitz.Matrix(2, 2)
    pix = page.get_pixmap(matrix=matrix, alpha=False)

    if target_ext == ".png":
        return pix.tobytes("png")
    if target_ext in {".jpg", ".jpeg"}:
        return pix.tobytes("jpg")
    if target_ext == ".webp":
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        buffer = io.BytesIO()
        image.save(buffer, "WEBP", quality=90, method=6)
        return buffer.getvalue()

    raise ValueError(f"Unsupported PDF page image target: {target_ext}")


def _write_pdf_page_images(input_path, output_dir, original_filename, target_ext):
    output_dir = pathlib.Path(output_dir)
    original_stem = pathlib.Path(original_filename).stem
    page_target_ext = ".png" if target_ext == ".zip" else target_ext
    file_ext = ".jpg" if page_target_ext == ".jpeg" else page_target_ext

    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")

        if doc.page_count == 1 and target_ext != ".zip":
            output_path = output_dir / f"{uuid.uuid4()}{file_ext}"
            output_path.write_bytes(_render_pdf_page_bytes(doc[0], page_target_ext))
            return ConversionOutput(
                path=str(output_path),
                filename=f"{original_stem}_page_1{file_ext}",
                media_type=PDF_MEDIA_TYPES[page_target_ext],
            )

        output_path = output_dir / f"{uuid.uuid4()}_{original_stem}_pages.zip"
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, page in enumerate(doc, start=1):
                archive.writestr(
                    f"{original_stem}/page_{index:04d}{file_ext}",
                    _render_pdf_page_bytes(page, page_target_ext),
                )

    if target_ext == ".zip":
        filename = get_pdf_download_name(original_filename, target_ext)
    else:
        filename = f"{original_stem}_{file_ext.lstrip('.')}_pages.zip"

    return ConversionOutput(
        path=str(output_path),
        filename=filename,
        media_type=PDF_MEDIA_TYPES[".zip"],
    )


def _write_compressed_pdf(input_path, output_path):
    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")
        doc.save(output_path, garbage=4, deflate=True, clean=True)
    return output_path


def _write_split_pdf_zip(input_path, output_path, original_filename):
    original_stem = pathlib.Path(original_filename).stem
    with fitz.open(input_path) as doc:
        if doc.page_count == 0:
            raise Exception("The PDF has no pages.")

        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for index in range(doc.page_count):
                single_page_pdf = fitz.open()
                single_page_pdf.insert_pdf(doc, from_page=index, to_page=index)
                pdf_bytes = single_page_pdf.tobytes(garbage=4, deflate=True)
                single_page_pdf.close()
                archive.writestr(f"{original_stem}/page_{index + 1:04d}.pdf", pdf_bytes)
    return output_path


async def process_pdf_conversion(input_path, output_dir, original_filename, target):
    output_dir = pathlib.Path(output_dir)

    if target in {".png", ".jpg", ".jpeg", ".webp", ".svg", ".zip"}:
        return await asyncio.to_thread(
            _write_pdf_page_images,
            input_path,
            output_dir,
            original_filename,
            target,
        )

    output_ext = ".pdf" if target == "compress" else ".zip" if target == "split" else target
    output_path = output_dir / f"{uuid.uuid4()}{output_ext}"

    writers = {
        ".docx": _write_pdf_docx,
        ".pptx": _write_pdf_pptx,
        ".xlsx": _write_pdf_xlsx,
        ".txt": _write_pdf_text_file,
        ".html": _write_pdf_html_file,
        "compress": _write_compressed_pdf,
        "split": lambda source, destination: _write_split_pdf_zip(source, destination, original_filename),
    }

    writer = writers.get(target)
    if not writer:
        raise ValueError(f"Unsupported PDF conversion target: {target}")

    await asyncio.to_thread(writer, input_path, str(output_path))
    return ConversionOutput(
        path=str(output_path),
        filename=get_pdf_download_name(original_filename, target),
        media_type=PDF_MEDIA_TYPES[target],
    )


def _write_merged_pdf(input_paths, output_path):
    merged = fitz.open()
    try:
        for input_path in input_paths:
            with fitz.open(input_path) as doc:
                if doc.page_count == 0:
                    raise Exception(f"The PDF has no pages: {pathlib.Path(input_path).name}")
                merged.insert_pdf(doc)
        if merged.page_count == 0:
            raise Exception("No pages were available to merge.")
        merged.save(output_path, garbage=4, deflate=True)
    finally:
        merged.close()
    return output_path


async def process_pdf_merge(input_paths, output_path):
    await asyncio.to_thread(_write_merged_pdf, input_paths, output_path)
    return output_path
